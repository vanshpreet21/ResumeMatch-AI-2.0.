
from pypdf import PdfReader
from docx import Document
from io import BytesIO


def extract_text(uploaded_file):
    filename = uploaded_file.name.lower()
    file_data = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        reader = PdfReader(BytesIO(file_data))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    elif filename.endswith(".docx"):
        document = Document(BytesIO(file_data))
        paragraphs = [
            p.text for p in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " ".join(cell.text for cell in row.cells)
                )

        return "\n".join(paragraphs)

    elif filename.endswith(".txt"):
        return file_data.decode("utf-8", errors="replace")

    else:
        raise ValueError("Unsupported file format")
